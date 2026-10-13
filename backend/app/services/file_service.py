"""Extract readable text from uploaded files."""
import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.utils.text_utils import clean_text


def extract_pdf_text(file_path: str) -> str:
    parts: list[str] = []
    with fitz.open(file_path) as pdf:
        for page in pdf:
            parts.append(page.get_text())
    return clean_text("\n\n".join(parts))


def extract_docx_text(file_path: str) -> str:
    document = DocxDocument(file_path)
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_text("\n".join(parts))


def extract_txt_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as fh:
        return clean_text(fh.read())
